"""Build ``Outfitter Guide.docx`` from ``Outfitter Guide.md``, formatted for print.

    uv run --no-project --with pypandoc_binary --with python-docx tools/build_guide_docx.py

Pandoc does the conversion; this script supplies the look. It writes a styled reference
document (fonts, heading colours, callouts, code, footer with page numbers) for pandoc to
take its styles from, then finishes what styles can't reach: the title page, table
borders and header shading, and image sizing.

The version on the title page is read from ``scripts/outfitter/__init__.py``.
"""
from __future__ import annotations

import re
import subprocess
import tempfile
from datetime import date
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parents[1]
SRC = REPO / "Outfitter Guide.md"
OUT = REPO / "Outfitter Guide.docx"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)       # headings, title, rules
ACCENT_HEX = "1F4E79"
INK = RGBColor(0x26, 0x26, 0x26)          # body text
MUTED = RGBColor(0x59, 0x59, 0x59)
TINT_HEX = "EAF1F8"                       # callout + table header fill
CODE_HEX = "F3F3F3"
RULE_HEX = "BFBFBF"
BODY_FONT, HEAD_FONT, CODE_FONT = "Calibri", "Calibri Light", "Consolas"
TEXT_WIDTH = Inches(6.5)                  # Letter, 1" margins


def _version() -> str:
    text = (REPO / "scripts" / "outfitter" / "__init__.py").read_text(encoding="utf-8")
    return re.search(r'__version__\s*=\s*"([^"]+)"', text).group(1)


# --- low-level XML helpers ----------------------------------------------------------

def _set_font(style_or_run, name: str) -> None:
    style_or_run.font.name = name
    rpr = style_or_run.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(attr), None)


def _shade(ppr_or_tcpr, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr_or_tcpr.append(shd)


def _border(parent, edges: dict[str, tuple[str, int, int]], tag: str = "w:pBdr") -> None:
    """``edges``: {"left": (hex colour, size in 1/8 pt, space in pt)}."""
    box = OxmlElement(tag)
    for edge, (color, size, space) in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), str(space))
        el.set(qn("w:color"), color)
        box.append(el)
    parent.append(box)


def _field(paragraph, instr: str) -> None:
    """Append a Word field (e.g. ``PAGE``) to ``paragraph``."""
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"),
                       ("end", None)):
        run = paragraph.add_run()
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def _style(doc, name: str, kind=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, kind)


# --- reference document: everything pandoc maps onto named styles -------------------

def build_reference(path: Path) -> None:
    pandoc = pypandoc.get_pandoc_path()
    with open(path, "wb") as fh:
        subprocess.run([pandoc, "--print-default-data-file", "reference.docx"],
                       stdout=fh, check=True)
    doc = Document(str(path))

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))

    normal = doc.styles["Normal"]
    _set_font(normal, BODY_FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name in ("Body Text", "First Paragraph", "Compact"):
        s = _style(doc, name)
        s.base_style = normal
        s.paragraph_format.space_before = Pt(0)
        s.paragraph_format.space_after = Pt(3 if name == "Compact" else 6)

    title = _style(doc, "Title")
    _set_font(title, HEAD_FONT)
    title.font.size, title.font.bold, title.font.color.rgb = Pt(36), False, ACCENT
    title.paragraph_format.space_before = Inches(2.2)
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.element.get_or_add_pPr().append(OxmlElement("w:pBdr"))  # drop any default rule
    sub = _style(doc, "Subtitle")
    _set_font(sub, HEAD_FONT)
    sub.font.size, sub.font.italic, sub.font.color.rgb = Pt(16), False, MUTED
    sub.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    sub.paragraph_format.space_after = Pt(30)
    date_style = _style(doc, "Date")
    date_style.font.size, date_style.font.color.rgb = Pt(11), MUTED
    date_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    sizes = {1: 22, 2: 15, 3: 12.5}
    for level, size in sizes.items():
        h = _style(doc, f"Heading {level}")
        _set_font(h, HEAD_FONT if level == 1 else BODY_FONT)
        h.font.size, h.font.color.rgb = Pt(size), ACCENT
        h.font.bold = level > 1
        h.font.italic = False
        pf = h.paragraph_format
        pf.space_before = Pt({1: 0, 2: 18, 3: 12}[level])
        pf.space_after = Pt({1: 14, 2: 6, 3: 4}[level])
        pf.keep_with_next = True
        pf.page_break_before = level == 1  # each numbered section starts a page
        if level == 1:
            _border(h.element.get_or_add_pPr(), {"bottom": (ACCENT_HEX, 8, 6)})

    toc_head = _style(doc, "TOC Heading")
    toc_head.base_style = doc.styles["Heading 1"]
    toc_head.paragraph_format.page_break_before = False

    callout = _style(doc, "Block Text")
    callout.base_style = normal
    callout.font.color.rgb = INK
    callout.font.italic = False
    cpf = callout.paragraph_format
    cpf.left_indent, cpf.right_indent = Inches(0.15), Inches(0.15)
    cpf.space_before, cpf.space_after = Pt(6), Pt(10)
    ppr = callout.element.get_or_add_pPr()
    _border(ppr, {"left": (ACCENT_HEX, 24, 8)})
    _shade(ppr, TINT_HEX)

    code = _style(doc, "Source Code")
    code.base_style = normal
    _set_font(code, CODE_FONT)
    code.font.size = Pt(9.5)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.space_before = code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.left_indent = Inches(0.15)
    ppr = code.element.get_or_add_pPr()
    _border(ppr, {e: (CODE_HEX, 4, 4) for e in ("top", "left", "bottom", "right")})
    _shade(ppr, CODE_HEX)
    verb = _style(doc, "Verbatim Char", WD_STYLE_TYPE.CHARACTER)
    _set_font(verb, CODE_FONT)
    verb.font.size = Pt(9.5)
    verb.font.color.rgb = RGBColor(0x8B, 0x1A, 0x1A)

    link = _style(doc, "Hyperlink", WD_STYLE_TYPE.CHARACTER)
    link.font.color.rgb = ACCENT
    cap = _style(doc, "Image Caption")
    cap.font.size, cap.font.italic, cap.font.color.rgb = Pt(9.5), True, MUTED
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    _style(doc, "Captioned Figure").paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style(doc, "Figure").paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer = sec.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run(f"Outfitter Guide  ·  v{_version()}  ·  ")
    run.font.size, run.font.color.rgb = Pt(9), MUTED
    _field(footer, "PAGE")
    for r in footer.runs[1:]:
        r.font.size, r.font.color.rgb = Pt(9), MUTED
    sec.different_first_page_header_footer = True  # no footer on the title page

    doc.save(str(path))


# --- finishing pass on pandoc's output ------------------------------------------------

def _style_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblpr = table._tbl.tblPr
    for old in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(old)
    _border(tblpr, {"top": (RULE_HEX, 4, 0), "bottom": (RULE_HEX, 4, 0),
                    "insideH": (RULE_HEX, 4, 0), "left": (RULE_HEX, 4, 0),
                    "right": (RULE_HEX, 4, 0), "insideV": (RULE_HEX, 4, 0)},
            tag="w:tblBorders")
    margins = OxmlElement("w:tblCellMar")
    for edge, twips in (("top", 50), ("bottom", 50), ("left", 100), ("right", 100)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(twips))
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    tblpr.append(margins)
    for r, row in enumerate(table.rows):
        if r == 0:  # repeat the header row across pages
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:tblHeader"))
        for cell in row.cells:
            if r == 0:
                _shade(cell._tc.get_or_add_tcPr(), TINT_HEX)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    if r == 0:
                        run.font.bold, run.font.color.rgb = True, ACCENT


def finish(path: Path, version: str) -> None:
    doc = Document(str(path))
    for table in doc.tables:
        _style_table(table)
    for shape in doc.inline_shapes:  # screenshots: fit the text width, keep the aspect
        if shape.width > TEXT_WIDTH * 0.9:
            ratio = (TEXT_WIDTH * 0.9) / shape.width
            shape.width, shape.height = int(shape.width * ratio), int(shape.height * ratio)
    # The TOC sits on its own page after the title page.
    for p in doc.paragraphs:
        if p.style.name == "Date":
            p.add_run().add_break(WD_BREAK.PAGE)
            break
    # Have Word fill in the table of contents (and page numbers) when the file opens.
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)
    doc.core_properties.title = "Outfitter Guide"
    doc.core_properties.subject = f"Outfitter v{version} user guide"
    doc.save(str(path))


def main() -> None:
    version = _version()
    with tempfile.TemporaryDirectory() as tmp:
        ref = Path(tmp) / "reference.docx"
        build_reference(ref)
        pypandoc.convert_file(
            str(SRC), "docx", format="gfm", outputfile=str(OUT),
            extra_args=[
                f"--reference-doc={ref}",
                f"--resource-path={REPO}",
                "--shift-heading-level-by=-1",   # '# Outfitter Guide' becomes the title
                "--toc", "--toc-depth=2",
                "-M", "toc-title=Contents",
                "-M", f"subtitle=Maya 2026 clothing tool  ·  Version {version}",
                "-M", f"date={date.today():%B %Y}",
            ])
    finish(OUT, version)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
